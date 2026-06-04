from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
import json
import sys


TEXT_EXTENSIONS = {'.md', '.markdown', '.txt'}
DOCX_EXTENSIONS = {'.docx'}
PDF_EXTENSIONS = {'.pdf'}


@dataclass
class ParsedDocumentResult:
    markdown: str
    plainText: str
    outline: list
    tables: list
    assets: list
    metadata: dict
    warnings: list


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def build_outline(markdown: str) -> list:
    outline = []
    for line in markdown.splitlines():
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            title = line[level:].strip()
            if title:
                outline.append({'level': level, 'title': title})
    return outline


def parse_text_document(path: Path) -> dict:
    content = path.read_text(encoding='utf-8')
    result = ParsedDocumentResult(
        markdown=content,
        plainText=content,
        outline=build_outline(content),
        tables=[],
        assets=[],
        metadata={
            'fileName': path.name,
            'extension': path.suffix.lower(),
            'parser': 'text',
        },
        warnings=[],
    )
    return asdict(result)


def parse_docx_document(path: Path, parser: str) -> dict:
    from docx import Document

    document = Document(str(path))
    markdown_lines = []
    plain_lines = []
    outline = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = str(getattr(paragraph.style, 'name', '') or '')
        if style_name.startswith('Heading '):
            try:
                level = int(style_name.split('Heading ')[1])
            except Exception:
                level = 1
            markdown_lines.append(f"{'#' * max(level, 1)} {text}")
            outline.append({'level': level, 'title': text})
        else:
            markdown_lines.append(text)

        plain_lines.append(text)

    markdown = '\n\n'.join(markdown_lines).strip() or '# 空文档'
    plain_text = '\n'.join(plain_lines).strip() or '空文档'
    result = ParsedDocumentResult(
        markdown=markdown,
        plainText=plain_text,
        outline=outline or build_outline(markdown),
        tables=[],
        assets=[],
        metadata={
            'fileName': path.name,
            'extension': path.suffix.lower(),
            'parser': parser,
            'hasOcr': False,
        },
        warnings=[],
    )
    return asdict(result)


def parse_pdf_document(path: Path, parser: str) -> dict:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    paragraphs = []

    for page_index in range(len(document)):
        page = document[page_index]
        text_page = page.get_textpage()
        text = text_page.get_text_range().strip()
        if text:
            paragraphs.append(text)
        text_page.close()
        page.close()

    document.close()

    markdown = '\n\n'.join(paragraphs).strip() or '# 空文档'
    plain_text = '\n\n'.join(paragraphs).strip() or '空文档'
    result = ParsedDocumentResult(
        markdown=markdown,
        plainText=plain_text,
        outline=build_outline(markdown),
        tables=[],
        assets=[],
        metadata={
            'fileName': path.name,
            'extension': path.suffix.lower(),
            'parser': parser,
            'hasOcr': False,
        },
        warnings=[],
    )
    return asdict(result)


def parse_document(file_path: str, parser: str = 'docling') -> dict:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f'file not found: {file_path}')

    extension = path.suffix.lower()

    if extension in TEXT_EXTENSIONS:
        return parse_text_document(path)

    if extension in DOCX_EXTENSIONS:
        return parse_docx_document(path, parser)

    if extension in PDF_EXTENSIONS:
        return parse_pdf_document(path, parser)

    result = ParsedDocumentResult(
        markdown='# 暂不支持\n\n当前仅支持 DOCX、PDF、Markdown 和 TXT 解析。',
        plainText='Unsupported document format.',
        outline=[{'level': 1, 'title': '暂不支持'}],
        tables=[],
        assets=[],
        metadata={
            'fileName': path.name,
            'extension': path.suffix.lower(),
            'parser': parser,
            'hasOcr': False,
        },
        warnings=[f'当前文件类型 {extension} 暂未接入解析能力。'],
    )
    return asdict(result)


if __name__ == '__main__':
    if len(sys.argv) >= 2:
        file_path = sys.argv[1]
        parser = sys.argv[2] if len(sys.argv) >= 3 else 'docling'
        print(json.dumps({
            'success': True,
            'generatedAt': now_iso(),
            'data': parse_document(file_path, parser=parser),
        }, ensure_ascii=False))
    else:
        sample_path = Path(__file__).resolve()
        print(json.dumps({
            'taskType': 'document-parse',
            'generatedAt': now_iso(),
            'sample': parse_document(str(sample_path), parser='docling'),
        }, ensure_ascii=False, indent=2))
